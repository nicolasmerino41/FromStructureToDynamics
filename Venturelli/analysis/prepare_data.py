"""Read original files without modifying them; extract a traceable analysis dataset."""
from pathlib import Path
import argparse,zipfile,xml.etree.ElementTree as ET,json,hashlib,csv
import numpy as np
import openpyxl

NAMES=['BH','CA','BU','PC','BO','BV','BT','EL','FP','CH','DP','ER']
S={'s':'http://www.sbml.org/sbml/level3/version1/core'}
def math_eval(node,env):
    tag=node.tag.split('}')[-1]
    if tag=='math':return math_eval(node[0],env)
    if tag=='ci':return env[node.text.strip()]
    if tag=='cn':return float(node.text)
    if tag=='apply':
        op=node[0].tag.split('}')[-1];v=[math_eval(x,env) for x in node[1:]]
        if op=='plus':return sum(v)
        if op=='times':return np.prod(v)
    raise ValueError('Unsupported MathML '+tag)

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--root',type=Path,default=Path(__file__).resolve().parents[1]);args=parser.parse_args()
    root=args.root;data=root/'Data';out=root/'analysis'/'processed';out.mkdir(parents=True,exist_ok=True)
    files=[{'name':p.name,'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(data.iterdir()) if p.is_file()]
    inventory={'source_files':files,'species_order':NAMES,'workbooks':{},'models':{}}
    (out/'species.txt').write_text('\n'.join(NAMES)+'\n')
    with zipfile.ZipFile(next(data.glob('*.zip'))) as archive:
        for tag in ['T1','T2','T3','T4']:
            xml=ET.fromstring(archive.read(f'twelveSpecies{tag}.xml'))
            pars={p.attrib['id']:float(p.attrib['value']) for p in xml.findall('.//s:listOfParameters/s:parameter',S)}
            u=np.array([pars['u'+i] for i in NAMES]);a=np.array([[pars['a'+i+j] for j in NAMES] for i in NAMES])
            assert np.all(np.diag(a)<0)
            # Verify orientation and formula by independently evaluating supplied MathML.
            state=np.linspace(.017,.11,12);env=pars|{'x'+n:v for n,v in zip(NAMES,state)}
            for reaction in xml.findall('.//s:reaction',S):
                product=reaction.find('s:listOfProducts/s:speciesReference',S).attrib['species'][1:]
                rate=math_eval(list(reaction.find('s:kineticLaw',S))[0],env)
                i=NAMES.index(product);assert np.isclose(rate,state[i]*(u[i]+a[i]@state),rtol=1e-12,atol=1e-14)
            initial={x.attrib['id']:float(x.attrib.get('initialConcentration',0)) for x in xml.findall('.//s:listOfSpecies/s:species',S)}
            x0=np.array([initial['x'+n] for n in NAMES])
            events=[]
            for event in xml.findall('.//s:event',S):
                ns={'m':'http://www.w3.org/1998/Math/MathML'}
                time=float(event.find('s:trigger',S).find('.//m:cn',ns).text)
                assignments=event.findall('.//s:eventAssignment',S)
                assert len(assignments)==12
                for assignment in assignments:
                    factor=float(assignment.find('.//m:cn',ns).text);assert factor==.05
                events.append({'time_hours':time,'fraction_remaining':.05})
            assert [x['time_hours'] for x in events]==[24.,48.]
            np.savetxt(out/f'{tag}_growth.csv',u,delimiter=',');np.savetxt(out/f'{tag}_interactions.csv',a,delimiter=',');np.savetxt(out/f'{tag}_initial.csv',x0,delimiter=',')
            inventory['models'][tag]={'events':events,'initial_abundance':x0.tolist(),'orientation':'row recipient, column source','mathml_verified':True}
    for path in sorted(data.glob('*.xlsx')):
        wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
        inventory['workbooks'][path.name]={s.title:{'rows':s.max_row,'columns':s.max_column,'description':[r[0] for r in s.iter_rows(values_only=True) if r[0] is not None] if s.title=='Description' else None} for s in wb}
        if 'MOESM4_' in path.name:
            groups={};key=None
            for (value,) in wb['Dataset EV3'].iter_rows(values_only=True):
                text=str(value).strip()
                if "'" in text:
                    key=text.strip(" '");groups[key]=[]
                else:
                    vals=[float(v) for v in text.split()]
                    assert len(vals)==12
                    groups[key].append(vals)
            assert 'NONE' in groups and all(len(g)==7 for g in groups.values())
            with (out/'observed_composition.csv').open('w',newline='') as f:
                w=csv.writer(f);w.writerow(['community','sample_index',*NAMES])
                for key,rows in groups.items():
                    for k,row in enumerate(rows):
                        assert abs(sum(row)-1)<.002
                        w.writerow([key,k,*row])
            np.savetxt(out/'observed_full.csv',np.array(groups['NONE']),delimiter=',')
    (out/'source_inventory.json').write_text(json.dumps(inventory,indent=2),encoding='utf-8')
    print('Extracted four models; verified all 48 reaction formulas and all dilution assignments. Original sources unchanged.')
if __name__=='__main__':main()

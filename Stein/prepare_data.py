"""Extract named parameters from the original, unmodified Stein Dataset S1."""
from pathlib import Path
import hashlib,json
import numpy as np
import openpyxl
root=Path(__file__).resolve().parent
src=root/'Data'/'stein2013_dataset_s1.xlsx'
out=root/'processed';out.mkdir(exist_ok=True)
w=openpyxl.load_workbook(src,read_only=True,data_only=True)
rows=list(w['MmuE'].values)
names=[r[0] for r in rows[2:13]]
assert names==list(rows[13][1:12]),'Parameter axes do not match'
A=np.array([r[1:12] for r in rows[2:13]],float)
r=np.array([v[12] for v in rows[2:13]],float)
eps=np.array([v[13] for v in rows[2:13]],float)
assert A.shape==(11,11) and np.all(np.diag(A)<0)
for filename,array in [('interactions',A),('growth',r),('susceptibilities',eps)]:
    np.savetxt(out/(filename+'.csv'),array,delimiter=',',fmt='%.17g')
(out/'species.txt').write_text('\n'.join(names)+'\n',encoding='utf8')
metadata={'dataset_url':'https://journals.plos.org/ploscompbiol/article/file?type=supplementary&id=10.1371/journal.pcbi.1003388.s001','paper_doi':'10.1371/journal.pcbi.1003388','sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'parameter_sheet':'MmuE','interaction_range':'B3:L13','growth_range':'M3:M13','susceptibility_range':'N3:N13','species':names,'time_unit':'day','abundance_unit':'10^11 DNA copies per cm^3 (source processed density proxy)','antibiotic_unit':'source scaled exposure, not a clinical dose','worksheets':{s.title:[s.max_row,s.max_column] for s in w}}
(out/'source_inventory.json').write_text(json.dumps(metadata,indent=2),encoding='utf8')
print('Extracted 11 groups, 121 interactions, growth and antibiotic susceptibilities.')
